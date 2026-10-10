import os
import docx

os.makedirs("sample_cvs", exist_ok=True)

# Sample 1: Strong Graphic Designer
doc1 = docx.Document()
doc1.add_heading("Alex Morgan", 0)
doc1.add_paragraph("Email: alex.morgan@example.com | Portfolio: https://behance.net/alexmorgan_design")
doc1.add_paragraph("Nationality: Canadian | Gender: Non-binary | Age: 29") # Demographic details to verify LLM ignores them
doc1.add_heading("Professional Summary", level=1)
doc1.add_paragraph("Senior Graphic Designer with 5 years of industry experience creating brand identities, UI designs, and digital marketing assets.")
doc1.add_heading("Skills", level=1)
doc1.add_paragraph("Adobe Photoshop, Adobe Illustrator, Figma, UI/UX Design, InDesign, Typography, Wireframing, Brand Identity")
doc1.add_heading("Experience", level=1)
doc1.add_paragraph("Lead Visual Designer - Creative Edge Studio (2021 - Present)\n- Designed complete design systems using Figma and Illustrator.\n- Managed digital asset creation in Photoshop.")
doc1.save("sample_cvs/Alex_Morgan_CV.docx")

# Sample 2: Junior Designer with Missing Experience
doc2 = docx.Document()
doc2.add_heading("Jordan Lee", 0)
doc2.add_paragraph("Email: jordan.l@example.com | Portfolio: https://github.com/jordanlee")
doc2.add_paragraph("Age: 22 | Male")
doc2.add_heading("Summary", level=1)
doc2.add_paragraph("Enthusiastic Junior Designer with 1 year of freelance graphic design experience.")
doc2.add_heading("Skills", level=1)
doc2.add_paragraph("Adobe Photoshop, Canva, Basic Illustrator, HTML/CSS")
doc2.add_heading("Experience", level=1)
doc2.add_paragraph("Freelance Designer (2023 - 2024)\n- Created social media banners using Photoshop.")
doc2.save("sample_cvs/Jordan_Lee_CV.docx")

print("Sample CVs generated in sample_cvs/")
